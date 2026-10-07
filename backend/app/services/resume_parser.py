import io
import os
from typing import Any, Dict, Optional, Tuple

import docx
import pymupdf

# Configuration limits
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB limit
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParserError(Exception):
    """Custom exception raised when resume parsing fails."""
    pass


def extract_text_from_pdf(content: bytes) -> Tuple[str, int]:
    """
    Extract text and page count from a PDF document in memory using PyMuPDF.

    Args:
        content: Raw bytes of the PDF file.

    Returns:
        Tuple of (extracted_text, page_count).
    """
    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ResumeParserError(f"Failed to read PDF file: {exc}") from exc

    page_count = len(doc)
    extracted_pages = []

    for page_num in range(page_count):
        page = doc.load_page(page_num)
        page_text = page.get_text()
        if page_text:
            extracted_pages.append(page_text.strip())

    doc.close()
    full_text = "\n\n".join(extracted_pages).strip()
    return full_text, page_count


def extract_text_from_docx(content: bytes) -> Tuple[str, Optional[int]]:
    """
    Extract text from a DOCX document in memory using python-docx.

    Args:
        content: Raw bytes of the DOCX file.

    Returns:
        Tuple of (extracted_text, None). (Pages are not fixed in Word documents).
    """
    try:
        doc_stream = io.BytesIO(content)
        doc = docx.Document(doc_stream)
    except Exception as exc:
        raise ResumeParserError(f"Failed to read DOCX file: {exc}") from exc

    text_parts = []

    # Extract text from standard paragraphs
    for paragraph in doc.paragraphs:
        cleaned_text = paragraph.text.strip()
        if cleaned_text:
            text_parts.append(cleaned_text)

    # Also extract text from tables (common for resumes formatted with tables)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate identical adjacent cell texts in merged cells
                deduped = []
                for cell_text in row_cells:
                    if not deduped or deduped[-1] != cell_text:
                        deduped.append(cell_text)
                text_parts.append(" | ".join(deduped))

    full_text = "\n".join(text_parts).strip()
    return full_text, None


def parse_resume(filename: str, content: bytes) -> Dict[str, Any]:
    """
    Validate and extract text from an uploaded resume file (PDF or DOCX).

    Args:
        filename: Original file name.
        content: Binary content of the uploaded file.

    Returns:
        Dictionary with metadata and extracted text.
    """
    if not content or len(content) == 0:
        raise ValueError("Uploaded file is empty.")

    if len(content) > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
        raise ValueError(
            f"File size ({len(content) / (1024 * 1024):.2f} MB) exceeds maximum allowed limit of {max_mb} MB."
        )

    _, ext = os.path.splitext(filename)
    normalized_ext = ext.lower().strip()

    if normalized_ext not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format '{ext}'. Only PDF (.pdf) and DOCX (.docx) files are supported."
        )

    file_type = normalized_ext.lstrip(".")

    if normalized_ext == ".pdf":
        text, pages = extract_text_from_pdf(content)
    elif normalized_ext == ".docx":
        text, pages = extract_text_from_docx(content)
    else:
        raise ValueError(f"Unsupported file extension: {normalized_ext}")

    return {
        "filename": filename,
        "file_type": file_type,
        "characters": len(text),
        "pages": pages,
        "text": text,
    }
